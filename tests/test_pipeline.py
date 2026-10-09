import csv
from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import load_workbook

from vehicle_data_quality.__main__ import demo, make_xlsx
from vehicle_data_quality.database import ROOT, connect, initialize, queries
from vehicle_data_quality.pipeline import COLUMNS, process_file, scan
from vehicle_data_quality.report import export_report, export_queries


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.database = self.directory / 'test.sqlite'
        initialize(self.database)

    def tearDown(self):
        self.temp.cleanup()

    def rows(self, sql):
        with connect(self.database) as db:
            return [dict(row) for row in db.execute(sql)]

    def write(self, name, records):
        path = self.directory / name
        with path.open('w', newline='', encoding='utf-8') as destination:
            writer = csv.DictWriter(destination, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(records)
        return path

    def record(self, **overrides):
        row = dict(zip(COLUMNS, ['TEST01', 'V001', 'Leasing', '123.45', '36', '10000', '2026-10-01']))
        return {**row, **overrides}

    def test_valid_import_counts_and_audit(self):
        result = process_file(self.database, ROOT / 'data/02_valid_updates.csv')
        self.assertEqual((result['inserted'], result['updated'], result['unchanged']), (2, 1, 1))
        self.assertEqual(len(self.rows('SELECT * FROM offers')), 6)
        changes = self.rows("SELECT * FROM changes WHERE offer_id='OFF001'")
        self.assertEqual(json.loads(changes[0]['old_json'])['price_cents'], 29900)
        self.assertEqual(json.loads(changes[0]['new_json'])['price_cents'], 34900)
        self.assertEqual(len(self.rows('SELECT * FROM import_rows')), 4)

    def test_five_quality_rules_reject_entire_file(self):
        before = self.rows('SELECT * FROM offers ORDER BY offer_id')
        result = process_file(self.database, ROOT / 'data/01_invalid_updates.csv')
        self.assertEqual(result['status'], 'rejected')
        codes = {row['code'] for row in self.rows('SELECT code FROM issues')}
        self.assertEqual(codes, {'REQUIRED', 'DUPLICATE', 'VALUE', 'FOREIGN_KEY', 'STALE_UPDATE'})
        self.assertEqual(self.rows('SELECT * FROM offers ORDER BY offer_id'), before)
        self.assertEqual(self.rows('SELECT * FROM changes'), [])

    def test_mixed_file_cannot_partially_apply(self):
        path = self.write('mixed.csv', [self.record(), self.record(offer_id='BAD', price_eur='-2')])
        self.assertEqual(process_file(self.database, path)['status'], 'rejected')
        self.assertFalse(any(row['offer_id'] == 'TEST01' for row in self.rows('SELECT * FROM offers')))

    def test_corrected_file_can_be_retried_and_duplicates_are_skipped(self):
        path = self.write('update.csv', [self.record(price_eur='')])
        self.assertEqual(process_file(self.database, path)['status'], 'rejected')
        self.write('update.csv', [self.record()])
        self.assertEqual(process_file(self.database, path)['inserted'], 1)
        self.assertEqual(process_file(self.database, path)['status'], 'skipped')
        self.assertEqual(len(self.rows('SELECT * FROM batches')), 2)
        self.assertEqual(len(self.rows('SELECT * FROM changes')), 1)

    def test_same_version_conflict_is_rejected(self):
        path = self.write('stale.csv', [self.record(offer_id='OFF001', updated_at='2026-09-01')])
        self.assertEqual(process_file(self.database, path)['status'], 'rejected')
        self.assertEqual(self.rows('SELECT code FROM issues')[0]['code'], 'STALE_UPDATE')

    def test_xlsx_round_trip_uses_typed_prices_and_dates(self):
        process_file(self.database, ROOT / 'data/02_valid_updates.csv')
        path = self.directory / 'updates.xlsx'
        make_xlsx(ROOT / 'data/02_valid_updates.csv', path)
        result = process_file(self.database, path)
        self.assertEqual((result['status'], result['unchanged']), ('accepted', 4))
        self.assertEqual(len(self.rows('SELECT * FROM changes')), 3)

    def test_rejects_xlsx_formulas(self):
        path = self.directory / 'formula.xlsx'
        make_xlsx(ROOT / 'data/02_valid_updates.csv', path)
        workbook = load_workbook(path)
        workbook.active['D2'] = '=1+1'
        workbook.save(path)
        workbook.close()
        result = process_file(self.database, path)
        self.assertEqual(result['status'], 'rejected')
        self.assertEqual(self.rows('SELECT code FROM issues')[0]['code'], 'FILE_SCHEMA')

    def test_malformed_csv_is_recorded(self):
        path = self.directory / 'broken.csv'
        path.write_text('wrong,columns\n1,2\n', encoding='utf-8')
        self.assertEqual(process_file(self.database, path)['status'], 'rejected')

    def test_malformed_xlsx_does_not_stop_inbox_processing(self):
        inbox = self.directory / 'broken-inbox'
        inbox.mkdir()
        (inbox / '01_broken.xlsx').write_bytes(b'not an Excel workbook')
        (inbox / '02_valid.csv').write_bytes((ROOT / 'data/02_valid_updates.csv').read_bytes())
        self.assertEqual([result['status'] for result in scan(self.database, inbox)], ['rejected', 'accepted'])
        self.assertEqual(self.rows('SELECT code FROM issues')[0]['code'], 'FILE_SCHEMA')

    def test_blank_middle_row_preserves_error_location(self):
        path = self.write('blank.csv', [self.record(), self.record(offer_id='BAD', price_eur='-1')])
        lines = path.read_text().splitlines()
        lines.insert(2, '')
        path.write_text('\n'.join(lines) + '\n')
        self.assertEqual(process_file(self.database, path)['status'], 'rejected')
        self.assertIn({'row_number': 4, 'code': 'VALUE'}, self.rows('SELECT row_number, code FROM issues'))

    def test_concurrent_identical_import_is_applied_once(self):
        path = ROOT / 'data/02_valid_updates.csv'
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: process_file(self.database, path), range(2)))
        self.assertEqual(sorted(result['status'] for result in results), ['accepted', 'skipped'])
        self.assertEqual(len(self.rows('SELECT * FROM batches')), 1)
        self.assertEqual(len(self.rows('SELECT * FROM changes')), 3)

    def test_nonfinite_fractional_and_out_of_range_values(self):
        for value in ['NaN', 'Infinity', '0', '123.456', '10000001']:
            with self.subTest(price=value):
                path = self.write(value.replace('.', '_') + '.csv', [self.record(price_eur=value)])
                self.assertEqual(process_file(self.database, path)['status'], 'rejected')

    def test_database_foreign_key_constraints_are_active(self):
        with connect(self.database) as db:
            with self.assertRaises(sqlite3.IntegrityError):
                db.execute("UPDATE offers SET vehicle_id='DOES_NOT_EXIST'")

    def test_scan_is_idempotent(self):
        inbox = self.directory / 'inbox'
        inbox.mkdir()
        (inbox / 'updates.csv').write_bytes((ROOT / 'data/02_valid_updates.csv').read_bytes())
        self.assertEqual(scan(self.database, inbox)[0]['status'], 'accepted')
        self.assertEqual(scan(self.database, inbox)[0]['status'], 'skipped')

    def test_eight_queries_and_service_aggregates(self):
        with connect(self.database) as db:
            analysis = queries(db)
        self.assertEqual(len(analysis), 8)
        self.assertEqual(analysis['07_service_by_category'][1][0], ['Fehlerklärung', 4, 33.0, 50.0])
        phase = analysis['08_service_by_phase'][1]
        self.assertIn(['manuell', 'Datenpflege', 2, 25.0], phase)
        self.assertIn(['automatisiert', 'Datenpflege', 2, 9.0], phase)

    def test_report_has_expected_values_tables_and_charts(self):
        process_file(self.database, ROOT / 'data/01_invalid_updates.csv')
        process_file(self.database, ROOT / 'data/02_valid_updates.csv')
        path = export_report(self.database, self.directory / 'report.xlsx')
        workbook = load_workbook(path)
        self.assertEqual(workbook['Übersicht']['B6'].value, 6)
        self.assertEqual(workbook['Übersicht']['B9'].value, 6)
        self.assertEqual(workbook['Übersicht']['B10'].value, 3)
        self.assertEqual(len(workbook['Übersicht']._charts), 1)
        self.assertEqual(len(workbook['Serviceanalyse']._charts), 1)
        self.assertEqual(len(workbook['Angebote'].tables), 1)
        self.assertEqual(workbook['Angebote']['E6'].data_type, 'n')
        for sheet in workbook:
            for row in sheet:
                self.assertFalse(any(cell.data_type == 'f' for cell in row))
        workbook.close()
        export_queries(self.database, self.directory / 'queries')
        self.assertEqual(len(list((self.directory / 'queries').glob('*.csv'))), 8)

    def test_full_demo(self):
        result = demo(self.directory / 'demo')
        self.assertEqual([run['status'] for run in result['results']], ['rejected', 'accepted', 'accepted', 'skipped'])
        self.assertTrue(Path(result['report']).is_file())
        with self.assertRaises(ValueError):
            demo(self.directory / 'demo')


if __name__ == '__main__':
    unittest.main()
