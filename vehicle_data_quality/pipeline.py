from collections import Counter
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .database import connect

COLUMNS = ('offer_id', 'vehicle_id', 'product', 'price_eur', 'term_months', 'annual_mileage', 'updated_at')
FIELDS = ('offer_id', 'vehicle_id', 'product', 'price_cents', 'term_months', 'annual_mileage', 'updated_at')
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


def text(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return '' if value is None else str(value).strip()


def read_rows(content, suffix):
    if suffix == '.csv':
        rows = list(csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True))
    elif suffix == '.xlsx':
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            rows = []
            for row in workbook.active.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError('Formeln sind in Importdateien nicht erlaubt.')
                rows.append([cell.value for cell in row])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError('Mehr als 10.000 Datensätze.')
        finally:
            workbook.close()
    else:
        raise ValueError('Nur CSV und XLSX werden unterstützt.')
    # Preserve physical row numbers in diagnostics; ignore only trailing empty rows.
    while rows and not any(text(value) for value in rows[-1]):
        rows.pop()
    if len(rows) < 2:
        raise ValueError('Datei enthält keine Datensätze.')
    header = [text(value) for value in rows[0]]
    if len(header) != len(COLUMNS) or set(header) != set(COLUMNS):
        raise ValueError('Spalten müssen genau dem dokumentierten Importschema entsprechen.')
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError('Mehr als 10.000 Datensätze.')
    rows[1:] = [row if row else [''] * len(header) for row in rows[1:]]
    if any(len(row) != len(header) for row in rows[1:]):
        raise ValueError('Mindestens eine Zeile hat eine falsche Spaltenanzahl.')
    return [dict(zip(header, [text(value) for value in row])) for row in rows[1:]]


def validate(rows, db):
    issues, candidates = [], []
    identifiers = Counter(row['offer_id'] for row in rows)
    vehicles = {row[0] for row in db.execute('SELECT id FROM vehicles')}
    existing = {row['offer_id']: dict(row) for row in db.execute('SELECT * FROM offers')}
    for number, row in enumerate(rows, start=2):
        start = len(issues)

        def issue(code, field, message):
            issues.append((number, code, field, message))

        missing = [column for column in COLUMNS if not row[column]]
        if missing:
            issue('REQUIRED', ', '.join(missing), 'Pflichtwert fehlt.')
        if row['offer_id'] and identifiers[row['offer_id']] > 1:
            issue('DUPLICATE', 'offer_id', 'Angebots-ID kommt mehrfach in der Datei vor.')
        if row['vehicle_id'] and row['vehicle_id'] not in vehicles:
            issue('FOREIGN_KEY', 'vehicle_id', 'Fahrzeug ist nicht im Stammdatenbestand vorhanden.')
        if missing:
            continue
        try:
            price = Decimal(row['price_eur'])
            term, mileage = Decimal(row['term_months']), Decimal(row['annual_mileage'])
            if not price.is_finite() or not 0 < price <= 10000000 or price * 100 != (price * 100).to_integral_value():
                raise ValueError('Preis muss positiv sein und höchstens zwei Nachkommastellen haben.')
            if not term.is_finite() or term != term.to_integral_value() or not 1 <= term <= 120:
                raise ValueError('Laufzeit muss eine ganze Zahl zwischen 1 und 120 sein.')
            if not mileage.is_finite() or mileage != mileage.to_integral_value() or not 1 <= mileage <= 200000:
                raise ValueError('Jahresfahrleistung muss zwischen 1 und 200.000 liegen.')
            if not re.fullmatch(r'[A-Z0-9_-]{1,40}', row['offer_id']):
                raise ValueError('Angebots-ID muss aus 1–40 Großbuchstaben, Ziffern, _ oder - bestehen.')
            if row['product'] not in ('Leasing', 'Finanzierung'):
                raise ValueError('Produkt muss Leasing oder Finanzierung sein.')
            source_date = date.fromisoformat(row['updated_at']).isoformat()
            candidate = dict(zip(FIELDS, [row['offer_id'], row['vehicle_id'], row['product'],
                             int(price * 100), int(term), int(mileage), source_date]))
        except (ValueError, InvalidOperation, OverflowError) as error:
            issue('VALUE', 'values', str(error))
            continue
        old = existing.get(row['offer_id'])
        if old and candidate != old and source_date <= old['updated_at']:
            issue('STALE_UPDATE', 'updated_at', 'Eine Änderung benötigt einen neueren Datenstand.')
        if len(issues) == start:
            candidates.append(candidate)
    return candidates, issues


def process_file(database, file):
    path = Path(file)
    # Read a bounded snapshot: the hashed bytes are exactly the bytes being validated.
    with path.open('rb') as source:
        content = source.read(MAX_FILE_BYTES + 1)
    digest = hashlib.sha256(content).hexdigest()
    now = datetime.now(timezone.utc).isoformat()
    with connect(database) as db:
        db.execute('BEGIN IMMEDIATE')
        old_batch = db.execute('SELECT * FROM batches WHERE sha256 = ?', (digest,)).fetchone()
        if old_batch:
            return {'status': 'skipped', 'batch_id': old_batch['id'], 'previous_status': old_batch['status']}
        try:
            if len(content) > MAX_FILE_BYTES:
                raise ValueError('Datei überschreitet 5 MiB.')
            rows = read_rows(content, path.suffix.lower())
            candidates, issues = validate(rows, db)
        except (ValueError, csv.Error, OSError, EOFError, BadZipFile, ParseError,
                InvalidFileException, KeyError) as error:
            rows, candidates, issues = [], [], [(0, 'FILE_SCHEMA', 'file', str(error))]
        status = 'rejected' if issues else 'accepted'
        cursor = db.execute('''INSERT INTO batches
            (file_name, sha256, status, started_at, completed_at, rows_total, rows_valid)
            VALUES (?, ?, ?, ?, ?, ?, ?)''',
            (path.name, digest, status, now, now, len(rows), len(candidates)))
        batch_id = cursor.lastrowid
        db.executemany('INSERT INTO import_rows VALUES (?, ?, ?)',
                       [(batch_id, number, json.dumps(row, ensure_ascii=False)) for number, row in enumerate(rows, 2)])
        db.executemany('INSERT INTO issues (batch_id, row_number, code, field, message) VALUES (?, ?, ?, ?, ?)',
                       [(batch_id, *issue) for issue in issues])
        counts = {'inserted': 0, 'updated': 0, 'unchanged': 0}
        if not issues:
            for candidate in candidates:
                old = db.execute('SELECT * FROM offers WHERE offer_id = ?', (candidate['offer_id'],)).fetchone()
                old = dict(old) if old else None
                if old == candidate:
                    counts['unchanged'] += 1
                    continue
                operation = 'update' if old else 'insert'
                counts['updated' if old else 'inserted'] += 1
                db.execute('''INSERT INTO offers VALUES (?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(offer_id) DO UPDATE SET vehicle_id=excluded.vehicle_id,
                    product=excluded.product, price_cents=excluded.price_cents,
                    term_months=excluded.term_months, annual_mileage=excluded.annual_mileage,
                    updated_at=excluded.updated_at''', tuple(candidate[field] for field in FIELDS))
                db.execute('''INSERT INTO changes (batch_id, offer_id, operation, old_json, new_json, changed_at)
                    VALUES (?, ?, ?, ?, ?, ?)''', (batch_id, candidate['offer_id'], operation,
                    json.dumps(old, ensure_ascii=False) if old else None,
                    json.dumps(candidate, ensure_ascii=False), now))
        db.execute('UPDATE batches SET inserted=?, updated=?, unchanged=?, completed_at=? WHERE id=?',
                   (*counts.values(), datetime.now(timezone.utc).isoformat(), batch_id))
        return {'status': status, 'batch_id': batch_id, 'rows': len(rows), 'valid_rows': len(candidates),
                'issues': len(issues), **counts}


def scan(database, inbox):
    # Producers should write a temporary file, then atomically rename it into the inbox.
    results = []
    for path in sorted(Path(inbox).iterdir()):
        if path.is_file() and path.suffix.lower() in ('.csv', '.xlsx') and not path.name.startswith('~$'):
            try:
                result = process_file(database, path)
            except OSError as error:
                result = {'status': 'error', 'message': str(error)}
            results.append({'file': path.name, **result})
    return results
