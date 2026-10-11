import csv
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .database import connect, queries

BLUE, INK = '214E7A', '1D2939'
LABELS = {
    'offer_id': 'Angebots-ID', 'brand': 'Marke', 'model': 'Modell', 'product': 'Produkt',
    'monthly_price_eur': 'Monatsrate (EUR)', 'term_months': 'Laufzeit (Monate)',
    'annual_mileage': 'Fahrleistung (km/Jahr)', 'updated_at': 'Datenstand',
    'batch_id': 'Import-ID', 'file_name': 'Datei', 'status': 'Status', 'rows_total': 'Zeilen',
    'rows_valid': 'Gültige Zeilen', 'inserted': 'Neu', 'updated': 'Geändert', 'unchanged': 'Unverändert',
    'started_at': 'Start (UTC)', 'completed_at': 'Ende (UTC)', 'row_number': 'Zeile',
    'code': 'Regel', 'field': 'Feld', 'message': 'Fehlerbeschreibung', 'operation': 'Vorgang',
    'old_price_eur': 'Alte Rate (EUR)', 'new_price_eur': 'Neue Rate (EUR)', 'changed_at': 'Änderung (UTC)',
    'category': 'Kategorie', 'requests': 'Anfragen', 'average_hours': 'Ø Dauer (Stunden)',
    'maximum_hours': 'Max. Dauer (Stunden)', 'phase': 'Beispielphase',
}


def put(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    # Treat source strings as literal text, including strings beginning with '='.
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name='Arial', size=11, color=INK)
    cell.alignment = Alignment(vertical='center', horizontal='left' if isinstance(value, str) else 'right')
    return cell


def prepare(workbook, name, title):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    put(sheet, 2, 1, title).font = Font(name='Arial', size=15, bold=True, color=INK)
    put(sheet, 3, 1, 'Synthetische Beispieldaten. Quelle: mitgelieferte CSV-Dateien und database.py.')
    return sheet


def table(sheet, headers, rows, name, start=5):
    for column, header in enumerate(headers, 1):
        cell = put(sheet, start, column, LABELS.get(header, header))
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for number, row in enumerate(rows, start + 1):
        for column, value in enumerate(row, 1):
            cell = put(sheet, number, column, value)
            if headers[column - 1].endswith('_eur'):
                cell.number_format = '#,##0.00'
            elif headers[column - 1].endswith('_hours'):
                cell.number_format = '0.0'
            if headers[column - 1] == 'updated_at' and isinstance(value, str):
                cell.value = datetime.strptime(value, '%Y-%m-%d').date()
                cell.number_format = 'yyyy-mm-dd'
    for column, header in enumerate(headers, 1):
        values = [LABELS.get(header, header)] + [str(row[column - 1] or '') for row in rows]
        width = min(58, max(16, max(map(len, values), default=16) + 3))
        sheet.column_dimensions[sheet.cell(start, column).column_letter].width = width
    for number, row in enumerate(rows, start + 1):
        lines = 1
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                cell = sheet.cell(number, column)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                width = sheet.column_dimensions[cell.column_letter].width
                lines = max(lines, (len(value) + int(width) - 1) // int(width))
        sheet.row_dimensions[number].height = max(22, lines * 17)
    sheet.row_dimensions[start].height = 26
    if rows:
        last = sheet.cell(start + len(rows), len(headers)).coordinate
        definition = Table(displayName=name, ref='A%d:%s' % (start, last))
        definition.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(definition)
    sheet.freeze_panes = 'A%d' % (start + 1)
    sheet.print_title_rows = '%d:%d' % (start, start)


def bar_chart(sheet, title, axis_title, column, start, end, anchor):
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = axis_title
    chart.add_data(Reference(sheet, min_col=column, min_row=start, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=start + 1, max_row=end))
    chart.legend = None
    chart.width, chart.height = 17, 9
    sheet.add_chart(chart, anchor)
    return chart


def export_report(database, destination):
    workbook = Workbook()
    workbook.remove(workbook.active)
    with connect(database) as db:
        db.execute('BEGIN')  # One consistent database snapshot for every worksheet.
        analysis = queries(db)
        summary = prepare(workbook, 'Übersicht', 'Vehicle Data Quality: Import und Datenqualität')
        summary.sheet_properties.tabColor = BLUE
        metrics = [
            ('Aktuelle Angebote', db.execute('SELECT COUNT(*) FROM offers').fetchone()[0]),
            ('Akzeptierte Dateien', db.execute("SELECT COUNT(*) FROM batches WHERE status='accepted'").fetchone()[0]),
            ('Abgelehnte Dateien', db.execute("SELECT COUNT(*) FROM batches WHERE status='rejected'").fetchone()[0]),
            ('Erkannte Fehler', db.execute('SELECT COUNT(*) FROM issues').fetchone()[0]),
            ('Protokollierte Änderungen', db.execute('SELECT COUNT(*) FROM changes').fetchone()[0]),
        ]
        table(summary, ['Kennzahl', 'Anzahl'], metrics, 'Kennzahlen')
        rule_labels = {'REQUIRED': 'Pflichtwerte', 'DUPLICATE': 'Doppelte IDs', 'VALUE': 'Wertebereiche',
                       'FOREIGN_KEY': 'Fahrzeugbezug', 'STALE_UPDATE': 'Datenstand', 'FILE_SCHEMA': 'Dateiformat'}
        errors = [(rule_labels.get(row[0], row[0]), row[1]) for row in analysis['04_errors_by_rule'][1]]
        table(summary, ['Qualitätsprüfung', 'Fehler'], errors, 'Qualitaetsregeln', start=14)
        summary.freeze_panes = None
        summary.column_dimensions['A'].width = 38
        summary.column_dimensions['B'].width = 14
        summary.column_dimensions['C'].width = 3
        if errors:
            chart = bar_chart(summary, 'Fehler nach Prüfregel', 'Anzahl', 2, 14, 14 + len(errors), 'D5')
            chart.type = 'bar'
            # In openpyxl BarChart, y_axis is the value axis even for horizontal bars.
            chart.y_axis.numFmt = '0'
            chart.y_axis.majorUnit = 1
            chart.y_axis.scaling.min = 0
            chart.y_axis.scaling.max = max(1, max(row[1] for row in errors))
        put(summary, 22, 1, 'Snapshot (UTC): ' + datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M'))
        put(summary, 23, 1, 'Aktualisierung: Bericht nach einem neuen Import erneut erzeugen.')

        for sheet_name, title, key, name in [
            ('Angebote', 'Aktueller Angebotsbestand', '01_current_offers', 'Angebote'),
            ('Importe', 'Importhistorie', '03_import_summary', 'Importe'),
            ('Änderungen', 'Änderungsnachweis mit alten und neuen Preisen', '05_change_history', 'Aenderungen'),
        ]:
            sheet = prepare(workbook, sheet_name, title)
            table(sheet, *analysis[key], name)
        sheet = prepare(workbook, 'Fehler', 'Fehlerprotokoll je Datei und Zeile')
        cursor = db.execute('''SELECT batch_id, row_number, code, field, message
                               FROM issues ORDER BY batch_id, row_number, id''')
        table(sheet, [column[0] for column in cursor.description], [list(row) for row in cursor], 'Fehler')
        sheet = prepare(workbook, 'Serviceanalyse', 'Bearbeitungszeiten synthetischer interner Anfragen')
        table(sheet, *analysis['07_service_by_category'], 'Servicekategorien')
        table(sheet, *analysis['08_service_by_phase'], 'Servicephasen', start=13)
        put(sheet, 22, 1, 'Phasen sind konstruierte Beispiele, kein Nachweis einer realen Zeitersparnis.')
        sheet.freeze_panes = 'A6'
        bar_chart(sheet, 'Mittlere Dauer nach Kategorie', 'Stunden', 3, 5, 8, 'F5')
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


def export_queries(database, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    with connect(database) as db:
        db.execute('BEGIN')
        for name, (headers, rows) in queries(db).items():
            with (directory / (name + '.csv')).open('w', encoding='utf-8-sig', newline='') as destination:
                writer = csv.writer(destination)
                writer.writerow(headers)
                # Prevent spreadsheet formula injection when CSVs are opened in Excel.
                writer.writerows([
                    ["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
                     for value in row] for row in rows
                ])
